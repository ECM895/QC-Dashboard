import pandas as pd
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import visualizations as vz
import datetime

status_color_map = {
    'A-Approved': '#16a34a',
    'B-Approved with Comments': '#22c55e',
    'C-Revise and Resubmit': '#d97706',
    'D-Rejected': '#b91c1c',
    'Closed': '#16a34a',
    'Open': '#b91c1c'
}

def generate_ppt(df, concrete_df, start_date, end_date):
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "QA/QC Construction Dashboard Report"
    subtitle.text = f"Report generated on {datetime.date.today()}\nFiltered from: {start_date} to {end_date}"
    
    # Generate Date Filtered df
    # Make sure Date is parsed if not already
    df['Date'] = pd.to_datetime(df['Date']).dt.date
    date_df = df[(df['Date'] >= start_date) & (df['Date'] <= end_date)]
    
    # --- SLIDE: Overall KPI Summary ---
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = txBox.text_frame
    tf.text = "Overall KPI Summary (Date Filtered)"
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    
    # Plot Overall Donut Chart
    overall_status = date_df['Status'].value_counts().reset_index()
    overall_status.columns = ['Status', 'Count']
    fig_donut = vz.plot_donut_chart(overall_status['Status'], overall_status['Count'], "OVERALL OBSERVATION STATUS", [status_color_map.get(s, '#808080') for s in overall_status['Status']])
    
    img_bytes = io.BytesIO()
    fig_donut.write_image(img_bytes, format='png', width=600, height=450, engine='kaleido')
    img_bytes.seek(0)
    slide.shapes.add_picture(img_bytes, Inches(2), Inches(2), height=Inches(4.5))
    
    # --- SLIDE: Metrics by Category ---
    slide = prs.slides.add_slide(blank_layout)
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = txBox.text_frame
    tf.text = "Metrics by Category (Date Filtered)"
    tf.paragraphs[0].font.size = Pt(28)
    
    cat_df = date_df.groupby(['Category', 'Status']).size().reset_index(name='Count')
    cat_totals = cat_df.groupby('Category')['Count'].transform('sum')
    cat_df['Percentage'] = (cat_df['Count'] / cat_totals) * 100
    
    fig_bar = vz.plot_100p_stacked_bar(cat_df, 'Category', 'Percentage', 'Status', "PROPORTIONAL STATUS BY CATEGORY", status_color_map)
    img_bytes = io.BytesIO()
    fig_bar.write_image(img_bytes, format='png', width=800, height=500, engine='kaleido')
    img_bytes.seek(0)
    slide.shapes.add_picture(img_bytes, Inches(1), Inches(1.5), width=Inches(8))
    
    # --- SLIDE: Concrete Productivity ---
    if concrete_df is not None and not concrete_df.empty:
        slide = prs.slides.add_slide(blank_layout)
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
        tf = txBox.text_frame
        tf.text = "Concrete Daily Productivity (Date Filtered)"
        tf.paragraphs[0].font.size = Pt(28)
        
        concrete_df['Date'] = pd.to_datetime(concrete_df['Date']).dt.date
        c_df = concrete_df[(concrete_df['Date'] >= start_date) & (concrete_df['Date'] <= end_date)].copy()
        if not c_df.empty:
            zone_el = c_df.groupby(['Zone', 'Element'])['Volume'].sum().reset_index()
            color_map = {'Columns': '#1e3a8a', 'Walls': '#3b82f6', 'Slab': '#93c5fd'}
            fig_c = vz.plot_grouped_bar_chart(zone_el, 'Zone', 'Volume', 'Element', "POURING BY ZONE & ELEMENT (m³)", color_map)
            img_bytes = io.BytesIO()
            fig_c.write_image(img_bytes, format='png', width=800, height=500, engine='kaleido')
            img_bytes.seek(0)
            slide.shapes.add_picture(img_bytes, Inches(1), Inches(1.5), width=Inches(8))

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()
