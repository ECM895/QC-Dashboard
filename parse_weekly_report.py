import pptx
import json
import os

prs = pptx.Presentation('auto_logs/Weekly Quality Report.pptx')
report_data = []

for idx, slide in enumerate(prs.slides):
    title = f"Slide {idx + 1}"
    if slide.shapes.title and slide.shapes.title.text:
        title = slide.shapes.title.text.strip().replace('\n', ' ')
    
    slide_entry = {
        'slide_number': idx + 1,
        'title': title,
        'text_blocks': [],
        'tables': []
    }
    
    for shape in slide.shapes:
        if shape.has_text_frame and shape != slide.shapes.title:
            txt = shape.text.strip()
            if txt and len(txt) > 2:
                slide_entry['text_blocks'].append(txt)
                
        if shape.has_table:
            t = shape.table
            table_data = []
            for r_idx in range(len(t.rows)):
                row_vals = [cell.text.strip().replace('\n', ' ') for cell in t.rows[r_idx].cells]
                table_data.append(row_vals)
            slide_entry['tables'].append(table_data)
            
    report_data.append(slide_entry)

with open('weekly_report_data.json', 'w') as f:
    json.dump(report_data, f, indent=2)

print(f"Extracted {len(report_data)} slides into weekly_report_data.json successfully!")
