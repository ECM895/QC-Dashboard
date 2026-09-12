import os
import pandas as pd
import numpy as np
import streamlit as st
import datetime
import glob

LOGS_DIR = "/Users/uzairahmad/Desktop/Logs"

def _generate_empty_data():
    """Generates an empty dataframe with correct columns to display 0s."""
    df = pd.DataFrame(columns=[
        'Date', 'Category', 'Reference No', 'Description', 
        'Status', 'Area', 'Contractor', 'Discipline'
    ])
    calibration_data = pd.DataFrame(columns=[
        'Equipment', 'Last Calibrated', 'Expiry Date', 'Status'
    ])
    return df, calibration_data

def process_uploaded_logs(_uploaded_files=None):
    """Reads Excel files from Streamlit uploader. If none exist, returns empty data."""
    is_mock = False
    dfs = []
    
    if _uploaded_files and len(_uploaded_files) > 0:
        # Exclude auxiliary trackers (Training, Lessons Learned, Concrete/Cast Insitu, Calibration, Templates)
        # so they do not pollute the Aconex master document register columns
        aux_keywords = ['learning event', 'training schedule', 'lessons_learned', 'best_practices', 'cast insitu', 'daily productivity', 'calibration_tracker', 'template_logs']
        log_files = [
            f for f in _uploaded_files 
            if not any(k in f.name.lower() for k in aux_keywords)
        ]
        
        for file in log_files:
            try:
                import io
                # Read all bytes into memory so we can re-read the stream multiple times
                file_bytes = io.BytesIO(file.read())
                
                temp_df = pd.read_excel(file_bytes, engine='calamine')
                
                # Auto-detect Custom NCR Log format
                if len(temp_df.columns) > 0 and "NON-CONFORMANCE REPORT" in str(temp_df.columns[0]):
                    # Dynamically find the header row to prevent offset crashes
                    header_idx = 15
                    for i, row in enumerate(temp_df.values):
                        if 'NCR  Document No' in str(row) or 'Sr.No' in str(row[0]):
                            header_idx = i + 1
                            break
                            
                    file_bytes.seek(0)
                    temp_df = pd.read_excel(file_bytes, skiprows=header_idx, engine='calamine')
                    
                    # Map NCR Log Columns to Dashboard Columns based on requested letters (B, J, L, R, S, U)
                    col_map = {}
                    for c in temp_df.columns:
                        col_name = str(c).strip()
                        if 'NCR  Document No' in col_name: col_map[c] = 'Reference No' # Column B
                        elif 'Location' in col_name: col_map[c] = 'Area'               # Column J
                        elif 'NCR Description' in col_name: col_map[c] = 'Description' # Column L
                        elif 'Date Responsed' in col_name: col_map[c] = 'Date'         # Column R
                        elif 'Current Status' in col_name: col_map[c] = 'Status'       # Column U
                        
                    temp_df = temp_df.rename(columns=col_map)
                    temp_df['Category'] = 'NCR'
                    temp_df['Contractor'] = 'Main Contractor'
                    
                    # Force NCR statuses to Open/Closed
                    if 'Status' in temp_df.columns:
                        temp_df = temp_df.dropna(subset=['Status'])
                        def map_ncr_status(val):
                            s = str(val).lower()
                            if 'close' in s: return 'Closed'
                            return 'Open'
                        temp_df['Status'] = temp_df['Status'].apply(map_ncr_status)

                # Auto-detect Aconex Exports
                elif len(temp_df.columns) > 0 and "In case any cell is highlighted" in str(temp_df.columns[0]):
                    
                    # Extract Category from the first 10 rows (User specified Row 5)
                    detected_category = None
                    for _, row in temp_df.head(10).iterrows():
                        row_str = str(row.values).lower()
                        if "work inspection request" in row_str or "type: wir" in row_str: detected_category = "WIR"
                        elif "material inspection request" in row_str or "type: mir" in row_str: detected_category = "MIR"
                        elif "material approval" in row_str or "type: mar" in row_str: detected_category = "MAR"
                        elif "method statement" in row_str or "method of statement" in row_str or "type: mst" in row_str: detected_category = "MST"
                        elif "inspection and test plan" in row_str or "inspection & test plan" in row_str or "inspection test plan" in row_str or "type: itp" in row_str or "'itp'" in row_str: detected_category = "ITP"
                        elif "shop drawing" in row_str or "type: shd" in row_str: detected_category = "SHD"
                        elif "non conformance" in row_str or "non-conformance" in row_str or "type: ncr" in row_str: detected_category = "NCR"
                        if detected_category:
                            break
                            
                    file_bytes.seek(0)
                    temp_df = pd.read_excel(file_bytes, skiprows=10, engine='calamine')
                    
                    # Map Aconex Columns to Dashboard Columns
                    col_map = {}
                    if 'Revision Date' in temp_df.columns: col_map['Revision Date'] = 'Date'
                    if 'Document No' in temp_df.columns: col_map['Document No'] = 'Reference No'
                    if 'Title' in temp_df.columns: col_map['Title'] = 'Description'
                    if 'Discipline' in temp_df.columns: col_map['Discipline'] = 'Area'
                    
                    temp_df = temp_df.rename(columns=col_map)
                    
                    # Parse Category from the 'Type' column if present (Multi-Type Support)
                    if 'Type' in temp_df.columns:
                        def map_aconex_type(row):
                            doc_no = str(row.get('Reference No', '')).strip()
                            s = str(row.get('Type', '')).lower()
                            if 'work inspection' in s: return 'WIR'
                            if 'material inspection' in s: return 'MIR'
                            if 'material approval' in s: return 'MAR'
                            if 'method statement' in s: return 'MST'
                            if 'test plan' in s or 'itp' in s: return 'ITP'
                            if 'shop drawing' in s or 'shd' in s: return 'SHD'
                            # Client Quality NCRs: Must start with or contain SOA-NCR-QL
                            if 'SOA-NCR-QL' in doc_no: return 'NCR'
                            if ('non conformance' in s or 'non-conformance' in s) and 'SOA-NCR-QL' in doc_no: return 'NCR'
                            return 'UNKNOWN'
                        temp_df['Category'] = temp_df.apply(map_aconex_type, axis=1)
                    else:
                        # Fallback: Infer Category from filename or Row 5
                        fname = file.name.upper()
                        if 'WIR' in fname: temp_df['Category'] = 'WIR'
                        elif 'MIR' in fname: temp_df['Category'] = 'MIR'
                        elif 'MAR' in fname: temp_df['Category'] = 'MAR'
                        elif 'MST' in fname: temp_df['Category'] = 'MST'
                        elif 'ITP' in fname: temp_df['Category'] = 'ITP'
                        elif 'SHD' in fname or 'SDH' in fname or 'SHOP DRAWING' in fname: temp_df['Category'] = 'SHD'
                        elif 'NCR' in fname: temp_df['Category'] = 'NCR'
                        elif detected_category: temp_df['Category'] = detected_category
                        else: temp_df['Category'] = 'UNKNOWN'
                    
                    # Map Status / Review Status to Exact Aconex Statuses
                    def map_status(row):
                        # For NCRs, check both 'Status' and 'Review Status'
                        raw_status = str(row.get('Status', '')).lower() if pd.notna(row.get('Status')) else ''
                        raw_review = str(row.get('Review Status', '')).lower() if pd.notna(row.get('Review Status')) else ''
                        combined_s = f"{raw_status} {raw_review}".strip()
                        cat = row.get('Category', '')
                        
                        if cat == 'NCR':
                            if any(x in combined_s for x in ['approved', 'closed', 'close', 'approve', 'pass']):
                                return 'Closed'
                            return 'Open'
                        else:
                            s = raw_review if raw_review and raw_review != 'nan' else raw_status
                            if not s or s == 'nan': return 'IGNORE'
                            if 'approved with comments' in s: return 'B-Approved with Comments'
                            if 'approved' in s and 'comments' not in s: return 'A-Approved'
                            if 'revise' in s or 'resubmit' in s: return 'C-Revise and Resubmit'
                            if 'reject' in s: return 'D-Rejected'
                            return 'IGNORE'
                            
                    temp_df['Status'] = temp_df.apply(map_status, axis=1)
                    # Filter out IGNORed statuses
                    temp_df = temp_df[temp_df['Status'] != 'IGNORE']
                        
                    temp_df['Contractor'] = 'Main Contractor' # Default for Aconex
                    
                # Normalize column names
                temp_df.columns = [str(c).strip() for c in temp_df.columns]
                # Try to parse dates
                if 'Date' in temp_df.columns:
                    temp_df['Date'] = pd.to_datetime(temp_df['Date'], errors='coerce').dt.date
                dfs.append(temp_df)
            except Exception as e:
                print(f"Error reading {file.name}: {e}")
                
        if len(dfs) > 0:
            df = pd.concat(dfs, ignore_index=True)
            if len(df) == 0:
                empty_df, mock_cal = _generate_empty_data()
                df = empty_df
                is_mock = True
            else:
                # Filter strictly to the 7 approved QA/QC Categories (remove all UNKNOWN)
                valid_qc_categories = ['WIR', 'MIR', 'MAR', 'MST', 'ITP', 'SHD', 'NCR']
                df = df[df['Category'].isin(valid_qc_categories)].copy()

                # For Shop Drawings (SHD), take into account only PDF drawings (exclude DWG files)
                is_shd = df['Category'] == 'SHD'
                is_pdf = (
                    df['File'].astype(str).str.lower().str.contains('pdf') | 
                    df['Reference No'].astype(str).str.upper().str.endswith('_PDF')
                )
                df = df[(~is_shd) | is_pdf].copy()

                # Ensure required columns exist to prevent crashes
                required_cols = ['Date', 'Category', 'Reference No', 'Description', 'Status', 'Area', 'Contractor']
                for col in required_cols:
                    if col not in df.columns:
                        df[col] = "N/A"
                        
                def extract_discipline(row):
                    text = str(row.get('Reference No', '')) + " " + str(row.get('Description', ''))
                    text = text.upper()
                    if ' CE ' in text or '-CE-' in text or 'CIVIL' in text: return 'Civil (CE)'
                    if ' AR ' in text or '-AR-' in text or 'ARCH' in text: return 'Arch (AR)'
                    if ' EL ' in text or '-EL-' in text or 'ELEC' in text: return 'Electrical (EL)'
                    if ' ME ' in text or '-ME-' in text or 'MECH' in text: return 'Mechanical (ME)'
                    if ' ST ' in text or '-ST-' in text or 'STEEL' in text: return 'Struc Steel (ST)'
                    return 'Other'
                    
                df['Discipline'] = df.apply(extract_discipline, axis=1)
                
                def extract_root_cause(row):
                    if row.get('Category') != 'NCR': return 'N/A'
                    text = str(row.get('Description', '')).lower()
                    if any(x in text for x in ['concrete', 'compressive', 'strength', 'cube']): return 'Concrete Strength Failure'
                    if any(x in text for x in ['rebar', 'formwork', 'alignment', 'cover', 'spacing']): return 'Formwork/Rebar Alignment'
                    if any(x in text for x in ['material', 'spec', 'approved', 'submittal', 'delivery']): return 'Material Specification'
                    if any(x in text for x in ['workmanship', 'finish', 'crack', 'honeycomb', 'damage', 'poor']): return 'Poor Workmanship'
                    if any(x in text for x in ['design', 'drawing', 'clash', 'dimension']): return 'Design/Drawing Issue'
                    return 'Other / Unclassified'
                    
                df['Root Cause'] = df.apply(extract_root_cause, axis=1)
        else:
            empty_df, mock_cal = _generate_empty_data()
            df = empty_df
            is_mock = True
    else:
        # Fallback if no files uploaded
        empty_df, mock_cal = _generate_empty_data()
        df = empty_df
        is_mock = True

    # 2. Calibration Data
    calibration_data = None
    if _uploaded_files:
        cal_file = next((f for f in _uploaded_files if f.name.endswith('calibration_tracker.xlsx')), None)
        if cal_file:
            try:
                import io
                cal_bytes = io.BytesIO(cal_file.read())
                calibration_data = pd.read_excel(cal_bytes)
            except:
                pass
                
    if calibration_data is None:
        _, calibration_data = _generate_empty_data()
        
    # 3. Concrete Data
    concrete_df = None
    if _uploaded_files:
        conc_file = next((f for f in _uploaded_files if 'concrete' in f.name.lower()), None)
        if conc_file:
            try:
                import io
                import datetime
                conc_bytes = io.BytesIO(conc_file.read())
                xls = pd.ExcelFile(conc_bytes)
                
                apportioned_data = []
                elements = ['Columns', 'External Wall', 'Internal Wall', 'Slab']
                
                for element in elements:
                    if element not in xls.sheet_names:
                        continue
                        
                    el_df = pd.read_excel(xls, sheet_name=element, header=None)
                    
                    m3_row_idx = None
                    for i, row in el_df.iterrows():
                        # The new files drop "(m3)" from the label, so make it resilient
                        row_str = str(row.values).lower()
                        if 'concrete quantity' in row_str or 'm3' in row_str:
                            m3_row_idx = i
                            break
                            
                    if m3_row_idx is None:
                        continue
                        
                    date_row_idx = None
                    for i in range(15):
                        has_date = False
                        for val in el_df.iloc[i, 4:10].values:
                            if isinstance(val, pd.Timestamp) or isinstance(val, datetime.datetime):
                                has_date = True
                                break
                        if has_date:
                            date_row_idx = i
                            break
                            
                    if date_row_idx is None:
                        continue
                        
                    dates = el_df.iloc[date_row_idx, 4:].values
                    
                    zone_data = el_df.iloc[date_row_idx+1:m3_row_idx, [2] + list(range(4, el_df.shape[1]))].copy()
                    zone_data.columns = ['Zone'] + list(dates)
                    
                    zone_data = zone_data.dropna(subset=['Zone'])
                    zone_data = zone_data[zone_data['Zone'].astype(str).str.contains('Zone', case=False)]
                    
                    valid_dates = [d for d in dates if pd.notna(d) and not isinstance(d, str)]
                    melted = pd.melt(zone_data, id_vars=['Zone'], value_vars=valid_dates)
                    melted.columns = ['Zone', 'Date', 'RawValue']
                    melted['RawValue'] = pd.to_numeric(melted['RawValue'], errors='coerce').fillna(0)
                    
                    m3_data = el_df.iloc[m3_row_idx, 4:].values
                    daily_m3 = pd.DataFrame({'Date': dates, 'DailyM3': m3_data})
                    daily_m3 = daily_m3[daily_m3['Date'].isin(valid_dates)]
                    daily_m3['DailyM3'] = pd.to_numeric(daily_m3['DailyM3'], errors='coerce').fillna(0)
                    
                    merged = pd.merge(melted, daily_m3, on='Date', how='left')
                    
                    daily_totals = merged.groupby('Date')['RawValue'].transform('sum')
                    merged['Volume'] = 0.0
                    mask = daily_totals > 0
                    merged.loc[mask, 'Volume'] = merged.loc[mask, 'DailyM3'] * (merged.loc[mask, 'RawValue'] / daily_totals[mask])
                    
                    merged['Element'] = element
                    apportioned_data.append(merged[['Zone', 'Date', 'Element', 'Volume']])
                    
                if apportioned_data:
                    final_df = pd.concat(apportioned_data, ignore_index=True)
                    final_df = final_df[final_df['Volume'] > 0]
                    final_df['Element'] = final_df['Element'].replace({'Internal Wall': 'Walls', 'External Wall': 'Walls'})
                    
                    # Fix Excel '1900-01-24' bug (when someone types '24' instead of a date)
                    final_df['Date'] = pd.to_datetime(final_df['Date'])
                    final_df = final_df[final_df['Date'].dt.year >= 2000]
                    final_df['Date'] = final_df['Date'].dt.date
                    
                    concrete_df = final_df
            except Exception as e:
                print(f"Error parsing concrete log: {e}")
                
    return df, calibration_data, concrete_df, is_mock

def filter_data(df, start_date, end_date, category):
    # Ensure dates are comparable as pandas Timestamps
    df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
    start_ts = pd.to_datetime(start_date)
    end_ts = pd.to_datetime(end_date)
    
    mask = (df['Date'] >= start_ts) & (df['Date'] <= end_ts)
    filtered = df.loc[mask]
    
    if category != "ALL":
        filtered = filtered[filtered['Category'] == category]
        
    return filtered

def parse_ncr_ppt(ppt_path_or_bytes):
    """Parses NCR tables from Open NCRs PowerPoint slide deck."""
    import pptx
    import re
    
    prs = pptx.Presentation(ppt_path_or_bytes)
    ppt_ncrs = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                for row_idx in range(1, len(table.rows)):
                    c = [cell.text.strip().replace('\n', ' ') for cell in table.rows[row_idx].cells]
                    if len(c) < 7: continue
                    raw_doc = c[1].replace('\x0b', ' ')
                    found = re.findall(r'SOA-NCR-QL-\d+|ECM-NCR-QL-\d+|NCR-QL-\d+|NCR-\d+', raw_doc)
                    key = found[0] if found else raw_doc
                    
                    # Area / Discipline from raw_doc if in brackets
                    area = ''
                    if '(' in raw_doc and ')' in raw_doc:
                        area = raw_doc[raw_doc.find('(')+1:raw_doc.find(')')]
                    
                    ppt_ncrs[key] = {
                        'Sr': c[0],
                        'DocRaw': raw_doc,
                        'Description': c[2],
                        'CAPA': c[3],
                        'IssueDate': c[4],
                        'DaysPassed': c[5],
                        'CurrentStatus': c[6],
                        'Area': area
                    }
    return ppt_ncrs

def get_ncr_master_data(all_submittals_df, ppt_path=None):
    """Reconciles cumulative Aconex NCR entries with Open NCRs PowerPoint tracker."""
    import re
    import datetime
    
    def normalize_zone(text, doc=''):
        full_text = f"{text} {doc}"
        m = re.search(r'\b(?:Zone|Z)\s*[-_:]?\s*([A-Za-z0-9]+(?:[- ][A-Za-z0-9]+)?)', full_text, re.IGNORECASE)
        if m:
            z = m.group(1).strip().upper()
            z = re.sub(r'\s+(WAS|EXEC|FOR|IN|AT|OF|WALL|RET|COL|MOCK)\b.*', '', z)
            if z.isdigit():
                return f"Zone {int(z)}"
            m_dig = re.match(r'^([0-9]+)[-_ ]', z)
            if m_dig:
                return f"Zone {int(m_dig.group(1))}"
            m_s = re.match(r'^[A-Z]\s+([0-9]+)', z)
            if m_s:
                return f"Zone {int(m_s.group(1))}"
            if 'B2' in z:
                return f"Zone {z}"
            return f"Zone {z}"
        return "Site-wide / General"

    # 1. Parse PPT if available
    ppt_ncrs = {}
    if ppt_path and os.path.exists(ppt_path):
        try:
            ppt_ncrs = parse_ncr_ppt(ppt_path)
        except Exception as e:
            print(f"Error parsing PPT {ppt_path}: {e}")
            
    # Also check auto_logs or logs folder for Open NCRs pptx
    if not ppt_ncrs:
        for search_dir in ["auto_logs", "logs", ".temp_uploads"]:
            if os.path.exists(search_dir):
                for fn in os.listdir(search_dir):
                    if fn.startswith("~$") or fn.startswith("."): continue
                    if fn.endswith('.pptx') and ('ncr' in fn.lower() or 'open' in fn.lower()):
                        try:
                            ppt_ncrs = parse_ncr_ppt(os.path.join(search_dir, fn))
                            if ppt_ncrs: break
                        except Exception as e:
                            print(f"Error reading PPT {fn}: {e}")
            if ppt_ncrs: break

    # 2. Filter NCRs from master submittal logs
    ncr_submittals = all_submittals_df[all_submittals_df['Category'] == 'NCR'].copy()
    if len(ncr_submittals) > 0:
        # Keep latest revision per Document / Reference No
        ncr_submittals = ncr_submittals.drop_duplicates(subset=['Reference No'], keep='last')

    ref_date = datetime.date(2026, 9, 12)
    reconciled = []
    seen_refs = set()

    # Match Aconex NCRs
    for _, r in ncr_submittals.iterrows():
        doc_no = str(r['Reference No']).strip()
        seen_refs.add(doc_no)
        
        # Match with PPT
        match_key = None
        for k in ppt_ncrs:
            if k in doc_no:
                match_key = k
                break
                
        p_data = ppt_ncrs.get(match_key, {}) if match_key else {}
        
        # Determine Date & Aging
        date_val = r['Date']
        if (pd.isna(date_val) or str(date_val) == 'NaT') and 'IssueDate' in p_data:
            date_val = p_data['IssueDate']
        date_dt = pd.to_datetime(date_val, errors='coerce')
        
        days_open = 0
        if pd.notna(date_dt):
            days_open = max(0, (ref_date - date_dt.date()).days)
            
        # Status determination:
        # If in PPT, check CurrentStatus
        ppt_status = p_data.get('CurrentStatus', '')
        aconex_status = str(r.get('Status', 'Open'))
        
        if ppt_status:
            if 'closed' in ppt_status.lower():
                final_status = 'Closed'
            else:
                final_status = 'Open'
        else:
            if any(x in aconex_status.lower() for x in ['closed', 'approved']):
                final_status = 'Closed'
            else:
                final_status = 'Open'
                
        # Aging category: Overdue if >= 60 days (2 months)
        if final_status == 'Open':
            aging_cat = 'Overdue (> 2 Months)' if days_open >= 60 else 'Active (< 2 Months)'
        else:
            aging_cat = 'Closed'
            
        capa = p_data.get('CAPA', '')
        if not capa:
            capa = 'Completed / Verified' if final_status == 'Closed' else 'Pending submission from contractor'
            
        status_note = ppt_status if ppt_status else ('Closed / Verified' if final_status == 'Closed' else 'Awaiting review / action')
        
        # Decode Discipline from numbering
        disc = r.get('Discipline', 'Civil (CE)')
        if p_data.get('Area'):
            disc = f"{p_data['Area']} — {disc}"
            
        desc_text = p_data.get('Description', r.get('Description', 'Quality Non-Conformance'))
        zone_val = normalize_zone(desc_text, doc_no)

        reconciled.append({
            'Document No': doc_no,
            'Description': desc_text,
            'Corrective Action': capa,
            'Issue Date': date_dt.date() if pd.notna(date_dt) else None,
            'Days Open': days_open,
            'Current Status': status_note,
            'Status': final_status,
            'Aging Category': aging_cat,
            'Discipline': disc,
            'Zone': zone_val,
            'Origin': 'Consultant (SOA)' if 'SOA-NCR' in doc_no else ('JV Internal (ECM)' if 'ECM-NCR' in doc_no else 'Other')
        })

    # Add any remaining PPT items that weren't in Aconex
    for k, p_data in ppt_ncrs.items():
        doc_raw = p_data.get('DocRaw', k)
        clean_doc = re.findall(r'[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+-NCR-[A-Z0-9]+-\d+', doc_raw)
        doc_no = clean_doc[0] if clean_doc else doc_raw
        if any(doc_no in s for s in seen_refs):
            continue
            
        date_dt = pd.to_datetime(p_data.get('IssueDate'), errors='coerce')
        days_open = 0
        if pd.notna(date_dt):
            days_open = max(0, (ref_date - date_dt.date()).days)
            
        ppt_desc = p_data.get('Description', 'Non-conformance recorded')
        zone_val = normalize_zone(ppt_desc, doc_no)
        ppt_status = p_data.get('Status', 'Open')
        final_status = 'Closed' if any(x in ppt_status.lower() for x in ['close', 'approv', 'pass']) else 'Open'
        aging_cat = 'Closed' if final_status == 'Closed' else ('Overdue (>14d)' if days_open > 14 else 'Within SLA (≤14d)')

        reconciled.append({
            'Document No': doc_no,
            'Description': ppt_desc,
            'Corrective Action': p_data.get('CAPA', 'Action defined'),
            'Issue Date': date_dt.date() if pd.notna(date_dt) else None,
            'Days Open': days_open,
            'Current Status': ppt_status,
            'Status': final_status,
            'Aging Category': aging_cat,
            'Discipline': p_data.get('Area', 'Civil / Arch'),
            'Zone': zone_val,
            'Origin': 'Consultant (SOA)' if 'SOA-NCR' in doc_no else 'JV Internal (ECM)'
        })

    master_df = pd.DataFrame(reconciled)
    if len(master_df) > 0:
        master_df = master_df.sort_values(by=['Days Open'], ascending=False)
    return master_df

def get_post_pour_data():
    import pandas as pd
    import random
    
    zones = ['Basement 1', 'Basement 2'] * 20
    defects = ['Honeycombing', 'Tie-rod holes', 'Crack', 'Uneven Surface']
    statuses = ['Open', 'Repaired', 'Inspected']
    contractors = ['Midmac', 'ECM']
    
    data = []
    for i in range(40):
        data.append({
            'Zone': zones[i],
            'Defect Type': random.choice(defects),
            'Status': random.choice(statuses),
            'Contractor': random.choice(contractors)
        })
        
    return pd.DataFrame(data)

@st.cache_data(ttl=300)
def get_training_data():
    """Loads ECM-JV Annual Quality and Functional Training records."""
    import os
    import pandas as pd
    
    qms_rows, func_rows = [], []
    train_file = None
    for d in ["auto_logs", "logs", "/Users/uzairahmad/Desktop/DG-2/Training", "/Users/uzairahmad/Desktop/Google/Quality Meeting"]:
        if os.path.exists(d):
            for fn in os.listdir(d):
                if fn.startswith("~$") or fn.startswith("."): continue
                if "training" in fn.lower() and (fn.endswith(".xlsx") or fn.endswith(".xls")):
                    train_file = os.path.join(d, fn)
                    break
        if train_file: break
        
    if train_file and os.path.exists(train_file):
        try:
            xls = pd.ExcelFile(train_file)
            month_names = {
                '2026-01-01 00:00:00': 'Jan 2026', '2026-02-28 00:00:00': 'Feb 2026',
                '2026-03-01 00:00:00': 'Mar 2026', '2026-04-01 00:00:00': 'Apr 2026',
                '2026-05-01 00:00:00': 'May 2026', '2026-06-01 00:00:00': 'Jun 2026',
                '2026-07-01 00:00:00': 'Jul 2026', '2026-08-01 00:00:00': 'Aug 2026',
                '2026-09-01 00:00:00': 'Sep 2026', '2026-10-01 00:00:00': 'Oct 2026',
                '2026-11-01 00:00:00': 'Nov 2026', '2026-12-01 00:00:00': 'Dec 2026'
            }
            for s in xls.sheet_names:
                is_qms = 'qms' in s.lower()
                is_func = 'functional' in s.lower()
                if not (is_qms or is_func): continue
                
                df_sheet = pd.read_excel(xls, sheet_name=s, header=None)
                m_row_idx = 13 if is_qms else 12
                w_row_idx = 14 if is_qms else 13
                start_row = w_row_idx + 2
                
                months_raw = df_sheet.iloc[m_row_idx].tolist() if m_row_idx < len(df_sheet) else []
                weeks_raw = df_sheet.iloc[w_row_idx].tolist() if w_row_idx < len(df_sheet) else []
                
                # Map column index to (Week, Month)
                curr_m = ''
                col_info = {}
                for c in range(7, len(weeks_raw)):
                    if c < len(months_raw) and pd.notna(months_raw[c]):
                        m_str = str(months_raw[c]).strip()
                        if m_str != '' and m_str.lower() != 'nan':
                            curr_m = month_names.get(m_str, m_str[:10])
                    w_str = str(weeks_raw[c]).strip() if pd.notna(weeks_raw[c]) else f"W{c}"
                    col_info[c] = (w_str, curr_m)
                    
                for i in range(start_row, len(df_sheet)):
                    row = df_sheet.iloc[i]
                    sn = row[0]
                    if pd.isna(sn) or str(sn).strip() == '' or str(sn).lower() == 'nan': continue
                    if str(sn).strip() == 'Week Number': continue
                    
                    pct = row[5]
                    pct_str = f"{float(pct)*100:.0f}%" if pd.notna(pct) and float(pct)<=1 else f"{float(pct):.0f}%" if pd.notna(pct) else "N/A"
                    
                    # Scan planned and completed weeks
                    completed_weeks = []
                    planned_events = []
                    for c in range(7, len(row)):
                        val = str(row[c]).strip().upper() if pd.notna(row[c]) else ''
                        if val in ['C', 'COMPLETED']:
                            w, m = col_info.get(c, ('?', ''))
                            completed_weeks.append(f"{w} ({m})" if m else w)
                        elif val in ['P', 'PLANNED', 'M', 'MOVED', 'L', 'OVERDUE']:
                            w, m = col_info.get(c, ('?', ''))
                            tag = "Overdue" if val in ['L', 'OVERDUE'] else ("Moved" if val in ['M', 'MOVED'] else "Planned")
                            planned_events.append((w, m, tag))
                            
                    if planned_events:
                        next_sched_val = f"📅 {planned_events[0][0]} • {planned_events[0][1]} ({planned_events[0][2]})"
                        status_tag = "Upcoming Scheduled"
                    elif completed_weeks:
                        next_sched_val = f"✅ Completed ({completed_weeks[-1]})"
                        status_tag = "Fully Completed"
                    else:
                        next_sched_val = "⏳ As Needed / Ad-hoc"
                        status_tag = "Routine"

                    entry = {
                        'S.N.': str(sn).strip(),
                        'Topic / Target Group': str(row[1]).strip() if pd.notna(row[1]) else ("General Quality" if is_qms else "Functional TBT"),
                        'Next Schedule': next_sched_val,
                        'Schedule Status': status_tag,
                        'Total Staff': row[2] if pd.notna(row[2]) else 0,
                        'Total Plan': row[3] if pd.notna(row[3]) else 0,
                        'Total Conducted': row[4] if pd.notna(row[4]) else 0,
                        '% Conducted': pct_str,
                        'Category': 'QMS Training' if is_qms else 'Functional / TBT'
                    }
                    if is_qms:
                        qms_rows.append(entry)
                    else:
                        func_rows.append(entry)
        except Exception as e:
            print(f"Error loading training data: {e}")
            
    df_qms = pd.DataFrame(qms_rows)
    df_func = pd.DataFrame(func_rows)
    return df_qms, df_func

@st.cache_data(ttl=300)
def get_lessons_learned_data():
    """Loads Diriyah Opera House Lessons Learned and Best Practices Register."""
    import os
    import pandas as pd
    
    ll_df, bp_df = pd.DataFrame(), pd.DataFrame()
    ll_file = None
    for d in ["auto_logs", "logs", "/Users/uzairahmad/Desktop/DG-2/Specs"]:
        if os.path.exists(d):
            for fn in os.listdir(d):
                if fn.startswith("~$") or fn.startswith("."): continue
                if "lesson" in fn.lower() and (fn.endswith(".xlsx") or fn.endswith(".xls")):
                    ll_file = os.path.join(d, fn)
                    break
        if ll_file: break
        
    if ll_file and os.path.exists(ll_file):
        try:
            xls = pd.ExcelFile(ll_file)
            for s in xls.sheet_names:
                if 'lesson' in s.lower():
                    df_raw = pd.read_excel(xls, sheet_name=s, header=None)
                    for i in range(len(df_raw)):
                        r = df_raw.iloc[i].dropna().tolist()
                        if any('Lesson Learned ID' in str(x) for x in r):
                            cols = [str(c).strip() for c in df_raw.iloc[i].tolist()]
                            clean = df_raw.iloc[i+1:].copy()
                            clean.columns = cols
                            ll_df = clean.dropna(subset=[cols[0]])
                            break
                elif 'best practice' in s.lower():
                    df_raw = pd.read_excel(xls, sheet_name=s, header=None)
                    for i in range(len(df_raw)):
                        r = df_raw.iloc[i].dropna().tolist()
                        if any('Best Practice ID' in str(x) for x in r):
                            cols = [str(c).strip() for c in df_raw.iloc[i].tolist()]
                            clean = df_raw.iloc[i+1:].copy()
                            clean.columns = cols
                            bp_df = clean.dropna(subset=[cols[0]])
                            break
        except Exception as e:
            print(f"Error loading lessons learned data: {e}")
            
    return ll_df, bp_df


